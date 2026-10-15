from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(11)

def h(text, level=1):
    doc.add_heading(text, level=level)

def para(text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    return p

def add_table(rows):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = 'Table Grid'
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = str(val)
            if i == 0:
                for r in cells[j].paragraphs[0].runs:
                    r.bold = True
    return t

title = doc.add_heading('PERANCANGAN SISTEM', 0)
sub = doc.add_paragraph('WEBSITE BUMDES DESA KALIBAKUNG')
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in sub.runs:
    r.bold = True

h('1. Use Case Diagram', 1)
para('Sistem memiliki lima aktor: Pengunjung (Warga), Calon Pengurus, Pengurus (Admin), Super Admin, dan Sistem Email.')
doc.add_picture('docs/diagrams/use-case.png', width=Inches(6.3))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
para('Gambar 1. Use Case Diagram', bold=True).alignment = WD_ALIGN_PARAGRAPH.CENTER

h('1.1 Ringkasan Use Case', 2)
rows = [['No', 'Use Case', 'Aktor']]
uc = [
 ['UC1', 'Melihat Beranda & Profil BUMDes', 'Pengunjung'],
 ['UC2', 'Melihat Unit Usaha', 'Pengunjung'],
 ['UC3', 'Melihat Laporan Keuangan Publik', 'Pengunjung'],
 ['UC4', 'Mengirim Aspirasi Warga', 'Pengunjung'],
 ['UC5', 'Mendaftar Akun (Register)', 'Calon Pengurus'],
 ['UC6', 'Verifikasi Email', 'Calon Pengurus + Sistem Email'],
 ['UC7', 'Login', 'Pengurus, Super Admin'],
 ['UC8', 'Mengelola Profil BUMDes', 'Pengurus'],
 ['UC9', 'Mengelola Unit Usaha', 'Pengurus'],
 ['UC10', 'Mengelola Laporan Keuangan', 'Pengurus'],
 ['UC11', 'Mengelola Aspirasi Warga', 'Pengurus'],
 ['UC12', 'Mengelola Notifikasi', 'Pengurus'],
 ['UC13', 'Mengelola Pengaturan Akun', 'Pengurus'],
 ['UC14', 'Menyetujui / Menolak Akun Pengurus', 'Super Admin'],
]
add_table(rows + uc)

h('2. Entity Relationship Diagram (ERD)', 1)
para('ERD memetakan relasi antar tabel. Basis data memakai PostgreSQL + Prisma; tipe data ditulis gaya MySQL agar seragam dengan format proposal.')
doc.add_picture('docs/diagrams/erd-full.png', width=Inches(6.3))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
para('Gambar 2. Tabel ERD Website BUMDes Desa Kalibakung', bold=True).alignment = WD_ALIGN_PARAGRAPH.CENTER

h('2.1 Penjelasan Relasi Antar Entitas', 2)
rows = [['No', 'Entitas', 'Relasi', 'Entitas', 'Kardinalitas', 'Keterangan']]
rels = [
 ['1', 'USERS', 'menginput', 'LAPORAN_KEUANGAN', '1 : M', 'Satu pengurus menginput banyak laporan; satu laporan diinput satu pengurus.'],
 ['2', 'BUMDES', 'memiliki', 'UNIT_USAHA', '1 : M', 'Satu BUMDes memiliki banyak unit usaha.'],
 ['3', 'BUMDES', 'menerbitkan', 'LAPORAN_KEUANGAN', '1 : M', 'Satu BUMDes menerbitkan banyak laporan keuangan bulanan.'],
 ['4', 'BUMDES', 'memiliki', 'PENGURUS', '1 : M', 'Satu BUMDes memiliki banyak pengurus (struktur organisasi).'],
]
add_table(rows + rels)
para('Catatan: ASPIRASI tidak memiliki foreign key ke users karena aspirasi dikirim pengunjung umum (warga) tanpa login. NOTIFIKASI bersifat log global untuk semua pengurus.')

h('3. Tabel ERD', 1)

tables = [
 ('Tabel 1. Perancangan Tabel Pengguna (users)', ['No', 'Nama Field', 'Type', 'Keterangan'], [
  ['1', 'id', 'Bigint(20)', 'Auto_increment, Primary Key'],
  ['2', 'nama', 'Varchar(255)', 'Not Null'],
  ['3', 'email', 'Varchar(255)', 'Not Null, Unique'],
  ['4', 'password', 'Varchar(255)', 'Not Null (hash bcrypt)'],
  ['5', 'jabatan', 'Varchar(255)', 'Null'],
  ['6', 'telepon', 'Varchar(255)', 'Null'],
  ['7', 'nik', 'Varchar(255)', 'Null'],
  ['8', 'unit_kerja', 'Varchar(255)', 'Null'],
  ['9', 'avatar', 'Varchar(255)', 'Null (nama file foto)'],
  ['10', 'role', 'Varchar(255)', 'SUPER_ADMIN / ADMIN / PENGURUS'],
  ['11', 'email_verified', 'Tinyint(1)', '0 = belum, 1 = terverifikasi'],
  ['12', 'approved', 'Tinyint(1)', '0 = menunggu, 1 = disetujui'],
  ['13', 'verify_token_hash', 'Varchar(255)', 'Null (hash token verifikasi)'],
  ['14', 'verify_token_expires', 'Timestamp', 'Null (kedaluwarsa 24 jam)'],
  ['15', 'created_at', 'Timestamp', 'Null'],
 ]),
 ('Tabel 2. Perancangan Tabel Data BUMDes (bumdes)', ['No', 'Nama Field', 'Type', 'Keterangan'], [
  ['1', 'id', 'Bigint(20)', 'Auto_increment, Primary Key'],
  ['2', 'nama', 'Varchar(255)', 'Not Null'],
  ['3', 'tahun_berdiri', 'Int(4)', 'Null'],
  ['4', 'dasar_hukum', 'Varchar(255)', 'Null (Perdes / SK)'],
  ['5', 'email', 'Varchar(255)', 'Null'],
  ['6', 'telepon', 'Varchar(255)', 'Null'],
  ['7', 'alamat', 'Text', 'Null'],
  ['8', 'deskripsi', 'Text', 'Null (slogan beranda publik)'],
  ['9', 'visi', 'Text', 'Null'],
  ['10', 'misi', 'Text', 'Null'],
  ['11', 'logo', 'Varchar(255)', 'Null (nama file logo)'],
  ['12', 'jam_operasional', 'Text', 'Null (format Hari | Jam)'],
 ]),
 ('Tabel 3. Perancangan Tabel Pengurus (pengurus)', ['No', 'Nama Field', 'Type', 'Keterangan'], [
  ['1', 'id', 'Bigint(20)', 'Auto_increment, Primary Key'],
  ['2', 'bumdes_id', 'Bigint(20)', 'Not Null, Foreign Key -> bumdes.id'],
  ['3', 'nama', 'Varchar(255)', 'Not Null'],
  ['4', 'jabatan', 'Varchar(255)', 'Not Null'],
  ['5', 'keterangan', 'Varchar(255)', 'Null (periode / No. SK)'],
  ['6', 'foto', 'Varchar(255)', 'Null (nama file foto)'],
  ['7', 'urutan', 'Int(10)', 'Default 0 (urutan tampil)'],
 ]),
 ('Tabel 4. Perancangan Tabel Unit Usaha (unit_usaha)', ['No', 'Nama Field', 'Type', 'Keterangan'], [
  ['1', 'id', 'Bigint(20)', 'Auto_increment, Primary Key'],
  ['2', 'bumdes_id', 'Bigint(20)', 'Not Null, Foreign Key -> bumdes.id'],
  ['3', 'nama', 'Varchar(255)', 'Not Null'],
  ['4', 'kategori', 'Varchar(255)', 'Not Null'],
  ['5', 'penanggung_jawab', 'Varchar(255)', 'Null'],
  ['6', 'kontak', 'Varchar(255)', 'Null'],
  ['7', 'deskripsi', 'Text', 'Null'],
  ['8', 'foto', 'Varchar(255)', 'Null (nama file foto)'],
  ['9', 'metrik_label', 'Varchar(255)', 'Null (mis. Omset Bulanan)'],
  ['10', 'metrik_nilai', 'Varchar(255)', 'Null (mis. Rp 18.250.000/bln)'],
  ['11', 'aktif', 'Tinyint(1)', 'Default 1 (tampil di publik)'],
  ['12', 'created_at', 'Timestamp', 'Null'],
 ]),
 ('Tabel 5. Perancangan Tabel Laporan Keuangan (laporan_keuangan)', ['No', 'Nama Field', 'Type', 'Keterangan'], [
  ['1', 'id', 'Bigint(20)', 'Auto_increment, Primary Key'],
  ['2', 'bumdes_id', 'Bigint(20)', 'Not Null, Foreign Key -> bumdes.id'],
  ['3', 'created_by', 'Bigint(20)', 'Null, Foreign Key -> users.id'],
  ['4', 'bulan', 'Int(2)', 'Not Null (1-12)'],
  ['5', 'tahun', 'Int(4)', 'Not Null'],
  ['6', 'total_pemasukan', 'Bigint(20)', 'Not Null'],
  ['7', 'total_pengeluaran', 'Bigint(20)', 'Not Null'],
  ['8', 'keterangan', 'Text', 'Null'],
  ['9', 'file_lpj', 'Varchar(255)', 'Null (PDF/Excel LPJ)'],
  ['10', 'status', 'Varchar(255)', 'DRAFT / PUBLIK'],
  ['11', 'created_at', 'Timestamp', 'Null'],
  ['12', 'updated_at', 'Timestamp', 'Null'],
 ]),
 ('Tabel 6. Perancangan Tabel Aspirasi (aspirasi)', ['No', 'Nama Field', 'Type', 'Keterangan'], [
  ['1', 'id', 'Bigint(20)', 'Auto_increment, Primary Key'],
  ['2', 'nama', 'Varchar(255)', 'Not Null'],
  ['3', 'kategori', 'Varchar(255)', 'Not Null'],
  ['4', 'whatsapp', 'Varchar(255)', 'Null (tidak ditampilkan publik)'],
  ['5', 'dusun', 'Varchar(255)', 'Null'],
  ['6', 'pesan', 'Text', 'Not Null'],
  ['7', 'status', 'Varchar(255)', 'MENUNGGU / DIPROSES / SELESAI'],
  ['8', 'catatan_internal', 'Text', 'Null (hanya admin)'],
  ['9', 'publik', 'Tinyint(1)', '0 = tidak, 1 = tampil di portal'],
  ['10', 'created_at', 'Timestamp', 'Null'],
 ]),
 ('Tabel 7. Perancangan Tabel Notifikasi (notifikasi)', ['No', 'Nama Field', 'Type', 'Keterangan'], [
  ['1', 'id', 'Bigint(20)', 'Auto_increment, Primary Key'],
  ['2', 'judul', 'Varchar(255)', 'Not Null'],
  ['3', 'pesan', 'Text', 'Not Null'],
  ['4', 'tipe', 'Varchar(255)', 'aspirasi / laporan / unit / sistem'],
  ['5', 'dibaca', 'Tinyint(1)', 'Default 0'],
  ['6', 'created_at', 'Timestamp', 'Null'],
 ]),
]
for t_title, header, data in tables:
    para(t_title, bold=True)
    add_table([header] + data)
    doc.add_paragraph()

h('4. Catatan Implementasi (PostgreSQL + Prisma)', 1)
rows = [['Gaya Proposal (MySQL)', 'Implementasi PostgreSQL']]
map_rows = [
 ['Bigint(20)', 'Int (auto-increment) / BigInt'],
 ['Varchar(255)', 'String (TEXT)'],
 ['Text', 'String (TEXT)'],
 ['Int(10)', 'Int'],
 ['Tinyint(1)', 'Boolean'],
 ['Timestamp', 'DateTime'],
 ['Auto_increment', '@id @default(autoincrement())'],
]
add_table(rows + map_rows)
para('Model Prisma yang sudah berjalan terdapat pada file prisma/schema.prisma.')

out = 'docs/PERANCANGAN-SISTEM-BUMDES-KALIBAKUNG.docx'
doc.save(out)
print('saved', out)
